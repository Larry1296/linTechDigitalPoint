"""Bulk product import from the stock spreadsheet (Excel, or CSV with the same headings).

One row is one sellable product or variant. Every row is checked before anything is saved, and a
committed import is all-or-nothing. Rules (see docs/product-import.md):

* Products are for internal stock management: imported products are never visible or orderable online.
* Categories are matched by name ignoring case and spacing; missing ones in a path such as
  "Mobile Accessories > Chargers" are created parent first. Where racks are in use, a category with no
  rack is reported for review and assigned only to a rack the person chooses.
* preferred_shelf must name one existing exact compartment. Shelves are never created from text.
* opening_quantity and opening_unit_cost are per sellable unit. Positive opening stock needs a shelf and a
  buying cost, and is created through the inventory service (lot, balance, movement, journal, audit).
* The SKU identifies a row. A row whose SKU or barcode is already in the system is skipped and changes nothing,
  so a file can be uploaded again, even after the products were edited or added by hand.
"""
import csv
import hashlib
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from uuid import uuid4

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import IntegrityError, transaction
from django.utils import timezone
from django.utils.text import slugify
from rest_framework.exceptions import ValidationError

from apps.core.models import AuditLog
from apps.tenancy.transactions import lock_business
from .models import Brand, Category, ProductImport, ProductImportRow, ProductVariant, shelves_in_use

MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024
# The template's headings, in order. Documented in docs/product-import.md.
COLUMNS = ["name", "category", "brand", "description", "product_type", "variant_name", "sku", "barcode", "selling_price", "minimum_price",
           "minimum_stock", "target_stock", "opening_quantity", "opening_unit_cost", "preferred_shelf"]
REQUIRED = ["name", "category", "sku", "selling_price"]
ONLINE_COLUMNS = {"ecommerce_visible", "online_orderable"}
TYPES = {"": "STOCK_ITEM", "stock_item": "STOCK_ITEM", "stock item": "STOCK_ITEM", "service": "SERVICE"}
MONEY_PLACES, QUANTITY_PLACES = 2, 3
LARGEST = Decimal("999999999999")  # fits the 14-digit money and quantity columns


def _key(text):
    """Names match regardless of case and spacing: " Mobile  accessories" is "mobile accessories"."""
    return " ".join(str(text or "").split()).casefold()


def _path_key(path):
    return " > ".join(_key(p) for p in str(path).split(">"))


def _heading(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _cell(value):
    """Spreadsheet cells as text. Excel stores long barcodes as numbers: 6161100000001.0 → "6161100000001"."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).upper()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _table(upload):
    name = (upload.name or "").lower()
    if upload.size > MAX_BYTES:
        raise ValidationError({"file": "The file is larger than 5 MB. Split it into smaller files."})
    raw = upload.read()
    if name.endswith(".xlsx") or raw[:2] == b"PK":
        from openpyxl import load_workbook
        try:
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheet = book["Products"] if "Products" in book.sheetnames else book.worksheets[0]
            return raw, [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        except Exception as exc:
            raise ValidationError({"file": "This Excel file could not be read. Save it again as .xlsx."}) from exc
    if name.endswith(".xls"):
        raise ValidationError({"file": "Old Excel (.xls) files aren't supported. Save it as .xlsx."})
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValidationError({"file": "The file's text encoding isn't recognised. Save it as .xlsx or CSV UTF-8."})
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return raw, [[_cell(v) for v in row] for row in csv.reader(io.StringIO(text), dialect)]


def _rows(table):
    """(row number, {column: text}) for every non-empty row. The headings must be the template's."""
    header_index = next((i for i, row in enumerate(table) if any(row)), None)
    if header_index is None:
        raise ValidationError({"file": "The file is empty."})
    headings = [_heading(h) for h in table[header_index]]
    named = [h for h in headings if h]
    online = sorted(ONLINE_COLUMNS.intersection(named))
    if online:
        raise ValidationError({"file": f"Remove the column{'s' if len(online) > 1 else ''} {', '.join(online)}. Imported products are "
                                       "always kept offline (not visible or orderable online); publish them from the product page later."})
    unknown = [h for h in named if h not in COLUMNS]
    missing = [c for c in REQUIRED if c not in named]
    repeated = sorted({h for h in named if named.count(h) > 1})
    if unknown or missing or repeated:
        problems = ([f"unknown heading{'s' if len(unknown) > 1 else ''}: {', '.join(unknown)}"] if unknown else []) + \
                   ([f"missing: {', '.join(missing)}"] if missing else []) + \
                   ([f"repeated: {', '.join(repeated)}"] if repeated else [])
        raise ValidationError({"file": f"Row {header_index + 1} must hold the template headings ({'; '.join(problems)}). "
                                       f"Expected: {', '.join(COLUMNS)}. Download the template to start from the exact headings."})
    rows = []
    for number, values in enumerate(table[header_index + 1:], start=header_index + 2):
        if not any(values):
            continue
        rows.append((number, {h: (values[i] if i < len(values) else "") for i, h in enumerate(headings) if h}))
    if not rows:
        raise ValidationError({"file": "The file has headings but no products."})
    if len(rows) > MAX_ROWS:
        raise ValidationError({"file": f"The file has {len(rows)} products; import at most {MAX_ROWS} at a time."})
    return rows


def _fingerprint(data):
    """The row as imported, ignoring spacing, letter case of names and number formatting (300 = 300.00)."""
    normal = {}
    for column in COLUMNS:
        text = " ".join(str(data.get(column, "")).split())
        if column == "category":
            text = _path_key(text)
        elif column in {"name", "brand", "product_type", "variant_name", "sku", "preferred_shelf"}:
            text = text.casefold()
        else:
            try:
                number = Decimal(text.replace(",", ""))
                if number.is_finite():
                    text = format(number.normalize(), "f")
            except InvalidOperation:
                pass
        if column == "variant_name" and text in ("", "standard"):
            text = ""
        normal[column] = text
    return hashlib.sha256(json.dumps(normal, sort_keys=True).encode()).hexdigest()


class _Row:
    def __init__(self, number, data):
        self.number, self.data = number, data
        self.errors, self.notes = [], []
        self.action = "create"
        self.category = self.shelf = None
        self.quantity = self.unit_cost = Decimal("0")

    def error(self, column, message):
        self.errors.append({"column": column, "message": message})

    def text(self, column):
        return " ".join(str(self.data.get(column, "")).split())

    def number_in(self, column, places, required=False, positive=False):
        """The column as a number, or None when blank or invalid (an error is recorded)."""
        text = self.text(column).replace(",", "")
        if not text:
            if required:
                self.error(column, f"{column} is required.")
            return None
        try:
            value = Decimal(text)
        except InvalidOperation:
            self.error(column, f"'{self.data[column]}' is not a number.")
            return None
        if not value.is_finite():
            self.error(column, f"'{self.data[column]}' is not a number.")
        elif value < 0:
            self.error(column, "cannot be negative.")
        elif positive and value == 0:
            self.error(column, "must be more than 0.")
        elif value > LARGEST:
            self.error(column, "is too large.")
        elif -value.as_tuple().exponent > places:
            self.error(column, f"use at most {places} decimal places.")
        else:
            return value
        return None


class _NewCategory:
    """A category this import would create."""

    def __init__(self, name, parent, parent_path):
        self.name, self.parent, self.pk = name, parent, None
        self.path = f"{parent_path} > {name}" if parent else name
        self.racks = set()


class _Categories:
    """Existing categories, the ones this file would add, and which racks hold each."""

    def __init__(self):
        from apps.inventory.models import ShelfStackCategory
        self.all = {c.pk: c for c in Category.objects.all()}
        self.children = {}
        for category in self.all.values():
            self.children.setdefault((category.parent_id, _key(category.name)), []).append(category)
        self.by_name = {}
        for category in self.all.values():
            self.by_name.setdefault(_key(category.name), []).append(category)
        self.racks = {}
        for assignment in ShelfStackCategory.objects.filter(stack__active=True, stack__zone__active=True):
            self.racks.setdefault(assignment.category_id, set()).add(assignment.stack_id)
        self.new = {}  # (parent identity, key) → _NewCategory, parents before children

    def path(self, node):
        if isinstance(node, _NewCategory):
            return node.path
        parts, seen = [], set()
        while node and node.pk not in seen:
            seen.add(node.pk)
            parts.append(node.name)
            node = self.all.get(node.parent_id)
        return " > ".join(reversed(parts))

    def lineage(self, node):
        while node is not None:
            yield node
            node = node.parent if isinstance(node, _NewCategory) else self.all.get(node.parent_id)

    def racks_for(self, node):
        """Racks that may hold this category: its own, or any parent's (a sub-category is stored on its parent's rack)."""
        racks = set()
        for item in self.lineage(node):
            racks |= item.racks if isinstance(item, _NewCategory) else self.racks.get(item.pk, set())
        return racks

    def assign(self, node, stack_id):
        if isinstance(node, _NewCategory):
            node.racks.add(stack_id)
        else:
            self.racks.setdefault(node.pk, set()).add(stack_id)

    def resolve(self, text, row):
        parts = [" ".join(p.split()) for p in text.split(">")]
        if any(not p for p in parts):
            row.error("category", f"'{text}' has an empty part. Write a path like Mobile Accessories > Chargers.")
            return None
        # A path is always written from the top: Parent > Child > Grandchild > ... The first name is a top-level
        # category and each ">" goes one level down, so every level is matched or created exactly where it is written.
        parent = None
        for depth, part in enumerate(parts):
            if len(part) > 120:
                row.error("category", f"'{part}' is longer than 120 characters.")
                return None
            key = _key(part)
            found = [] if isinstance(parent, _NewCategory) else self.children.get((parent.pk if parent else None, key), [])
            if not found and depth == 0:
                elsewhere = sorted(self.path(c) for c in self.by_name.get(key, []))
                if elsewhere:
                    # Created as written, but say so: the person may have meant the existing sub-category.
                    row.notes.append(f"'{part}' is not a top-level category, so a new top-level '{part}' will be created. "
                                     f"It already exists as {'; '.join(elsewhere)}: write that full path if you meant it.")
            if len(found) > 1:
                row.error("category", f"'{self.path(found[0])}' exists {len(found)} times. Rename or hide the duplicates under Categories first.")
                return None
            if found:
                node = found[0]
                if not node.active:
                    row.error("category", f"'{self.path(node)}' is hidden. Reactivate it under Categories or use another category.")
                    return None
            else:
                identity = (id(parent) if isinstance(parent, _NewCategory) else (parent.pk if parent else None), key)
                node = self.new.get(identity)
                if node is None:
                    node = self.new[identity] = _NewCategory(part, parent, self.path(parent) if parent else "")
            parent = node
        return parent


class _Shelves:
    """Exact compartments (a shelf on an active level of an active rack), found by code or printed label."""

    def __init__(self):
        from apps.inventory.models import Shelf
        self.by_text = {}
        for shelf in Shelf.objects.filter(is_virtual=False).select_related("zone", "level__stack"):
            for label in {shelf.code, shelf.physical_label}:
                if label and label.strip():
                    self.by_text.setdefault(_key(label), {})[shelf.pk] = shelf

    @staticmethod
    def usable(shelf):
        return (shelf.active and shelf.zone.active and shelf.level_id is not None
                and shelf.level.active and shelf.level.stack.active)

    @staticmethod
    def label(shelf):
        where = f"{shelf.level.stack.display_name}, level {shelf.level.level_number}" if shelf.level_id else shelf.zone.name
        return f"{shelf.physical_label or shelf.code} ({where})"

    def find(self, text, row):
        matches = list(self.by_text.get(_key(text), {}).values())
        if not matches:
            row.error("preferred_shelf", f"'{text}' is not an existing shelf compartment. Use a compartment code or label from the "
                                         "template's Shelves list. Shelves are never created by an import.")
            return None
        usable = [s for s in matches if self.usable(s)]
        if len(usable) > 1:
            row.error("preferred_shelf", f"'{text}' is ambiguous: it matches {', '.join(sorted(self.label(s) for s in usable))}. "
                                         "Use the compartment's unique code.")
            return None
        if not usable:
            shelf = matches[0]
            reason = "is not an exact compartment on a rack level" if shelf.level_id is None else "is archived, or on an archived rack or area"
            row.error("preferred_shelf", f"'{text}' {reason}. Choose an active compartment.")
            return None
        return usable[0]


class _Check:
    def __init__(self, parsed, rack_choices):
        from apps.inventory.models import ShelfStack
        self.categories = _Categories()
        self.shelves = _Shelves()
        self.racks_required = shelves_in_use()
        self.stacks = {s.pk: s for s in ShelfStack.objects.filter(active=True, zone__active=True).select_related("zone").order_by("display_name")}
        self.brands = {_key(b.name): b for b in Brand.objects.all()}
        self.variants = {v.sku.casefold(): v for v in ProductVariant.objects.select_related("product")}
        self.barcodes = {v.barcode: v for v in self.variants.values() if v.barcode}
        self.imported = {}
        for record in ProductImportRow.objects.filter(variant__isnull=False).select_related("batch"):
            self.imported.setdefault(record.variant_id, []).append(record)
        self.rows = [_Row(n, d) for n, d in parsed]
        self.file_errors = []
        for row in self.rows:
            self.check(row)
        self.duplicates()
        self.apply_rack_choices(rack_choices)
        self.rack_issues = self.check_racks()

    def check(self, row):
        name, sku = row.text("name"), row.text("sku")
        row.sku, row.name, row.barcode = sku, name, row.text("barcode")
        row.fingerprint = _fingerprint(row.data)
        existing = self.variants.get(sku.casefold()) if sku else None
        if existing:
            # An import never changes existing products: a SKU already in the system is skipped, not an error.
            row.action, row.variant = "skip", existing
            earlier = [r for r in self.imported.get(existing.pk, []) if r.fingerprint == row.fingerprint]
            if earlier:
                row.notes.append(f"Already imported in {earlier[0].batch.reference} (row {earlier[0].row_number}): skipped, nothing changes.")
            else:
                row.notes.append(f"SKU {existing.sku} already exists ({existing.product.name}, {existing.name}): skipped, nothing changes. "
                                 "Edit it on its product page to change it.")
            return
        taken = self.barcodes.get(row.barcode) if row.barcode else None
        if taken:
            row.action, row.variant = "skip", taken
            row.notes.append(f"Barcode {row.barcode} already belongs to {taken.product.name} (SKU {taken.sku}): skipped, nothing changes.")
            return
        if not name:
            row.error("name", "name is required.")
        elif len(name) > 180:
            row.error("name", "is longer than 180 characters.")
        if not sku:
            row.error("sku", "sku is required.")
        elif len(sku) > 64:
            row.error("sku", "is longer than 64 characters.")
        barcode = row.barcode
        if len(barcode) > 64:
            row.error("barcode", "is longer than 64 characters.")
        kind = TYPES.get(row.text("product_type").casefold().replace("-", "_"))
        if kind is None:
            row.error("product_type", f"'{row.data['product_type']}' is not a product type. Write STOCK_ITEM or SERVICE, or leave it empty.")
            kind = "STOCK_ITEM"
        row.product_type = kind
        row.variant_name = row.text("variant_name") or "Standard"
        if len(row.variant_name) > 120:
            row.error("variant_name", "is longer than 120 characters.")
        brand = row.text("brand")
        row.brand = brand
        if len(brand) > 120:
            row.error("brand", "is longer than 120 characters.")
        row.description = str(row.data.get("description", "")).strip()
        row.selling_price = row.number_in("selling_price", MONEY_PLACES, required=True, positive=True)
        # The bargaining floor: the lowest price a cashier may accept. Empty means no bargaining without a manager.
        row.minimum_price = row.number_in("minimum_price", MONEY_PLACES) or Decimal("0")
        if row.selling_price and row.minimum_price > row.selling_price:
            row.error("minimum_price", f"({row.minimum_price}) cannot be higher than selling_price ({row.selling_price}).")
        minimum = row.number_in("minimum_stock", QUANTITY_PLACES) or Decimal("0")
        target = row.number_in("target_stock", QUANTITY_PLACES) or Decimal("0")
        if target and target < minimum:
            row.error("target_stock", f"({target}) cannot be lower than minimum_stock ({minimum}).")
        row.minimum_stock, row.target_stock = minimum, target
        quantity = row.number_in("opening_quantity", QUANTITY_PLACES)
        cost = row.number_in("opening_unit_cost", MONEY_PLACES)
        shelf_text = row.text("preferred_shelf")
        category = row.text("category")
        if not category:
            row.error("category", "category is required.")
        else:
            row.category = self.categories.resolve(category, row)
        if kind == "SERVICE":
            # A service is sold, never stocked: no opening stock, cost or shelf.
            for column, value in (("opening_quantity", quantity), ("opening_unit_cost", cost)):
                if value:
                    row.error(column, "must be empty for a SERVICE: services have no stock.")
            if shelf_text:
                row.error("preferred_shelf", "must be empty for a SERVICE: services are not kept on a shelf.")
        else:
            if shelf_text:
                row.shelf = self.shelves.find(shelf_text, row)
            if quantity:
                row.quantity = quantity
                if not shelf_text:
                    row.error("preferred_shelf", f"is required for opening stock of {quantity}: name the compartment where it is kept.")
                if cost is None and not any(e["column"] == "opening_unit_cost" for e in row.errors):
                    row.error("opening_unit_cost", f"is required for opening stock of {quantity}: enter the buying cost of one unit.")
                elif cost == 0:
                    row.error("opening_unit_cost", "must be more than 0 for opening stock: enter the buying cost of one unit.")
                elif cost:
                    row.unit_cost = cost
            elif cost:
                row.notes.append("opening_unit_cost is ignored because opening_quantity is empty.")

    def duplicates(self):
        """The same new SKU or barcode twice in the file is an error on every row using it. Skipped rows change nothing."""
        for column in ("sku", "barcode"):
            seen = {}
            for row in self.rows:
                if row.action != "skip" and getattr(row, column):
                    seen.setdefault(getattr(row, column).casefold(), []).append(row)
            for rows in seen.values():
                for row in rows if len(rows) > 1 else []:
                    others = ", ".join(str(r.number) for r in rows if r is not row)
                    row.error(column, f"{getattr(row, column)} is also used on row{'s' if len(rows) > 2 else ''} {others}. Each {column} must be unique in the file.")

    def apply_rack_choices(self, choices):
        """Racks the person chose in the review for categories of this file: never guessed."""
        self.rack_choices = []
        # A rack may be chosen for a row's category or any of its parents (sub-categories are stored on the parent's rack).
        nodes = {_path_key(self.categories.path(node)): node
                 for r in self.rows if r.category is not None for node in self.categories.lineage(r.category)}
        for path, stack_id in (choices or {}).items():
            if stack_id in (None, ""):
                continue
            node = nodes.get(_path_key(path))
            try:
                stack = self.stacks.get(int(stack_id))
            except (TypeError, ValueError):
                stack = None
            if node is None:
                self.file_errors.append(f"Rack choice for '{path}': that category is not in this file.")
            elif stack is None:
                self.file_errors.append(f"Rack choice for '{path}': choose an active rack.")
            else:
                self.categories.assign(node, stack.pk)
                self.rack_choices.append((node, stack))

    def check_racks(self):
        """Categories that need a rack, and compartments on a rack that does not hold the row's category."""
        issues = {}
        for row in self.rows:
            if row.action == "skip" or row.category is None:
                continue
            racks = self.categories.racks_for(row.category)
            path = self.categories.path(row.category)
            if self.racks_required and not racks:
                # Nothing in this category's family has a rack, so the rack is asked for once, for the top category: every
                # sub-category under it then uses that rack, and no parent is left without one.
                top = list(self.categories.lineage(row.category))[-1]
                top_path = self.categories.path(top)
                issue = issues.setdefault(top_path, {"category": top_path, "new": isinstance(top, _NewCategory), "rows": [], "suggested_racks": []})
                issue["rows"].append(row.number)
                if row.shelf and row.shelf.level.stack_id not in issue["suggested_racks"]:
                    issue["suggested_racks"].append(row.shelf.level.stack_id)
                row.error("category", f"'{path}' is not on a rack yet, so it can't hold products. Choose a rack for '{top_path}' in the review "
                                      "(its sub-categories use the same rack), or assign it under Shelf management, and check again.")
            elif row.shelf and row.shelf.level.stack_id not in racks:
                stack = row.shelf.level.stack
                row.error("preferred_shelf", f"{self.shelves.label(row.shelf)} is on rack '{stack.display_name}', which does not hold "
                                             f"category '{path}'. Use a compartment on a rack for this category, or add the category to that rack.")
                issue = issues.setdefault(path, {"category": path, "new": isinstance(row.category, _NewCategory), "rows": [], "suggested_racks": []})
                issue["rows"].append(row.number)
                if stack.pk not in issue["suggested_racks"]:
                    issue["suggested_racks"].append(stack.pk)
        return list(issues.values())

    def result(self):
        rows = []
        valid = [r for r in self.rows if not r.errors]
        for row in self.rows:
            value = (row.quantity * row.unit_cost).quantize(Decimal("0.01"))
            rows.append({
                "row": row.number, "name": row.name, "sku": row.sku, "variant_name": getattr(row, "variant_name", ""),
                "product_type": getattr(row, "product_type", ""), "action": row.action,
                "category": self.categories.path(row.category) if row.category is not None else row.text("category"),
                "shelf": self.shelves.label(row.shelf) if row.shelf else "",
                "opening_quantity": str(row.quantity), "opening_unit_cost": str(row.unit_cost), "opening_value": str(value),
                "errors": row.errors, "notes": row.notes,
            })
        creating = [r for r in valid if r.action == "create"]
        new_brands = sorted({r.brand for r in creating if r.brand and _key(r.brand) not in self.brands}, key=str.casefold)
        summary = {
            "rows": len(self.rows), "create": len(creating), "skip": sum(1 for r in valid if r.action == "skip"),
            "errors": sum(1 for r in self.rows if r.errors), "error_count": sum(len(r.errors) for r in self.rows) + len(self.file_errors),
            "file_errors": self.file_errors,
            "stock_items": sum(1 for r in creating if r.product_type == "STOCK_ITEM"),
            "services": sum(1 for r in creating if r.product_type == "SERVICE"),
            "with_opening_stock": sum(1 for r in creating if r.quantity > 0),
            "total_quantity": str(sum((r.quantity for r in creating), Decimal("0"))),
            "opening_value": str(sum(((r.quantity * r.unit_cost).quantize(Decimal("0.01")) for r in creating), Decimal("0.00"))),
            "new_categories": [n.path for n in self.categories.new.values()],
            "new_brands": new_brands,
            "rack_issues": self.rack_issues,
            "rack_choices": [{"category": self.categories.path(n), "rack_id": s.pk, "rack": s.display_name} for n, s in self.rack_choices],
            "racks": [{"id": s.pk, "name": s.display_name, "area": s.zone.name} for s in self.stacks.values()] if self.rack_issues or self.rack_choices else [],
            "online": False,
        }
        return {"rows": rows, "summary": summary, "committed": False}


def _previous_import(digest):
    earlier = ProductImport.objects.filter(sha256=digest).select_related("created_by").first()
    if not earlier:
        return None
    return {"reference": earlier.reference, "created_at": earlier.created_at.isoformat(),
            "by": earlier.created_by.get_username() if earlier.created_by else ""}


def _unique_slug(name):
    base = slugify(name)[:40].strip("-") or "category"
    slug, suffix = base, 2
    while Category.objects.filter(slug=slug).exists():
        slug, suffix = f"{base}-{suffix}", suffix + 1
    return slug


def _save(check, user, upload, digest):
    """Writes every change of a clean check. Called inside the import's transaction."""
    from apps.inventory.models import ShelfStackCategory
    from apps.inventory.services import create_product_with_opening_stock
    reference = f"IMPORT-{timezone.now():%y%m%d}-{uuid4().hex[:6].upper()}"
    result = check.result()
    summary = result["summary"]
    batch = ProductImport.objects.create(reference=reference, file_name=(upload.name or "")[:255], sha256=digest, created_by=user,
                                         summary={k: summary[k] for k in ("rows", "create", "skip", "total_quantity", "opening_value", "new_categories", "new_brands")})
    # Only a clean file is saved, so every new category belongs to a row being created. Parents come first.
    for node in check.categories.new.values():
        category = Category.objects.create(name=node.name, slug=_unique_slug(node.name), parent_id=node.parent.pk if node.parent else None)
        node.pk = category.pk
        AuditLog.objects.create(action="CATEGORY_CREATED", object_type="Category", object_id=str(category.pk), user=user,
                                after={"name": category.name, "path": node.path, "import": reference})
    for node, stack in check.rack_choices:
        _, created = ShelfStackCategory.objects.get_or_create(stack=stack, category_id=node.pk, defaults={"is_primary": False})
        if created:
            AuditLog.objects.create(action="RACK_CATEGORY_ASSIGNED", object_type="ShelfStack", object_id=str(stack.pk), user=user,
                                    after={"category": check.categories.path(node), "rack": stack.display_name, "import": reference})
    brands = dict(check.brands)
    for row, public in zip(check.rows, result["rows"]):
        if row.action != "create":
            continue
        brand = None
        if row.brand:
            brand = brands.get(_key(row.brand)) or Brand.objects.create(name=row.brand)
            brands[_key(row.brand)] = brand
        category = row.category if isinstance(row.category, Category) else Category.objects.get(pk=row.category.pk)
        try:
            product, variant = create_product_with_opening_stock(
                user=user,
                # Internal stock only: never shown or orderable online, so no photo is needed.
                product_data={"name": row.name, "category": category, "brand": brand, "product_type": row.product_type,
                              "description": row.description, "ecommerce_visible": False, "online_orderable": False},
                variant_data={"name": row.variant_name, "sku": row.sku, "barcode": row.barcode or None, "selling_price": row.selling_price, "minimum_price": row.minimum_price,
                              "minimum_stock": row.minimum_stock, "target_stock": row.target_stock},
                preferred_shelf=row.shelf, opening_quantity=row.quantity, opening_unit_cost=row.unit_cost, opening_reference=reference)
        except (ValidationError, DjangoValidationError, IntegrityError) as exc:
            detail = getattr(exc, "detail", None) or getattr(exc, "messages", None) or str(exc)
            raise ValidationError({"file": f"Row {row.number} ({row.sku}) could not be saved: {_flatten(detail)} Nothing was imported."}) from exc
        ProductImportRow.objects.create(batch=batch, row_number=row.number, variant=variant, sku=row.sku, fingerprint=row.fingerprint)
        public.update(product_id=product.pk, variant_id=variant.pk)
    AuditLog.objects.create(action="PRODUCTS_IMPORTED", object_type="ProductImport", object_id=reference, user=user,
                            after={"sha256": digest, "file": upload.name, "reference": reference, **batch.summary})
    result.update(committed=True, reference=reference)
    return result


def _flatten(detail):
    if isinstance(detail, dict):
        return " ".join(f"{k}: {_flatten(v)}" for k, v in detail.items())
    if isinstance(detail, (list, tuple)):
        return " ".join(_flatten(v) for v in detail)
    return str(detail)


def import_products(upload, user, commit=False, rack_choices=None):
    """Check (and with commit=True, save) every row of a product spreadsheet. rack_choices maps a category path to a rack id."""
    if rack_choices and not _may_assign_racks(user):
        raise ValidationError({"rack_choices": "Assigning a category to a rack needs the Shelf management permission. Ask a manager."})
    raw, table = _table(upload)
    parsed = _rows(table)
    digest = hashlib.sha256(raw).hexdigest()
    if not commit:
        result = _Check(parsed, rack_choices).result()
        result["summary"]["previous_import"] = _previous_import(digest)
        return result
    with transaction.atomic():
        # Checked again under the business lock, so nothing can change between check and save.
        lock_business()
        check = _Check(parsed, rack_choices)
        result = check.result()
        summary = result["summary"]
        summary["previous_import"] = _previous_import(digest)
        if summary["error_count"]:
            return result
        if not summary["create"]:
            # Every row was imported before: nothing to save, and nothing wrong.
            result.update(committed=True, reference="")
            return result
        result = _save(check, user, upload, digest)
        result["summary"]["previous_import"] = summary["previous_import"]
        return result


def _may_assign_racks(user):
    from apps.tenancy.permissions import has_business_permission
    return has_business_permission(user, "inventory.change_shelfstack")
