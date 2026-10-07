"""The Excel template for importing products. Its headings are exactly product_import.COLUMNS, one for each field
of the Add Product form, and each row is saved just as that form would save it.

The first sheet ("Products") is the one imported. It is an Excel table with only headings, so a new row typed
under the last one carries the dropdowns and formats with it, and the same file can be filled, uploaded and
added to again (rows already imported are skipped). Examples live on their own sheet. Dropdowns list this
business's own categories, brands and shelf compartments; new categories (at any depth) and brands may still be
typed in and are created on import. Shelves are never created: they must already exist. A business that does not
use the Shelves & locations tool gets no preferred_shelf column, as the form shows no shelf picker.
"""
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .product_import import COLUMNS, MAX_ROWS, REQUIRED

ROWS = MAX_ROWS  # rows prepared with dropdowns and formats: as many as one import takes
TEXT, MONEY, QUANTITY, WHOLE = "@", "#,##0.00", "General", "0"
# heading → (width, number format, help text)
FIELDS = {
    "name": (34, TEXT, "Required. What the item is called, e.g. Oraimo Type-C cable. One row per sellable product or variant."),
    "category": (40, TEXT, "Required. The full path from the top-level category, each level separated by >: "
                           "Parent > Child > Grandchild > ... e.g. Mobile Accessories > Chargers > Car Chargers. "
                           "Any number of levels. Every level not yet in the system is created, parent first, exactly where it is written. "
                           "Letter case and extra spaces are ignored when matching."),
    "brand": (18, TEXT, "Optional. Choose from the list or type a new brand."),
    "description": (36, TEXT, "Optional short description."),
    "product_type": (16, TEXT, "STOCK_ITEM (kept in stock and sold) or SERVICE (work you do, no stock). Empty means STOCK_ITEM."),
    "variant_name": (20, TEXT, "Optional size, colour or model, e.g. 1m, Black, 45W. Empty means Standard "
                               "(or 'Pack of 10 clips' when sold in packs)."),
    "pack_size": (12, WHOLE, "How is it sold? Empty or 1: one at a time (a charger, a mouse). 2 or more: in packs, that many pieces for "
                             "one price (paper 2 sheets for 5, cable clips 10 for 20). Then the price, the till and minimum/target stock "
                             "count packs. Stock items only."),
    "piece_name": (14, TEXT, "Only with a pack_size of 2 or more: what one piece is called, e.g. sheet, clip. Empty means piece."),
    "sku": (22, TEXT, "Optional. Your unique code for this product or variant. Leave empty and it is made from the brand, name and "
                      "variant, as the Add Product form suggests (Oraimo · Type-C cable · 1m → ORAIMO-TCC-1M). Must not be used by "
                      "another product."),
    "barcode": (18, TEXT, "Optional, stock items only. The number under the barcode. Must be unique. Typed as text so Excel keeps leading zeros."),
    "selling_price": (14, MONEY, "Required. Price of one sellable unit (one pack when sold in packs), more than 0."),
    "minimum_price": (14, MONEY, "Optional bargaining guide: the lowest price staff should agree, per unit. Cannot be higher than "
                                 "selling_price. Empty or 0 means no bargaining."),
    "service_cost": (14, MONEY, "SERVICE rows only, optional: what doing the service costs you each time. Empty means 0."),
    "minimum_stock": (14, QUANTITY, "Optional, stock items. You are warned to restock when stock falls to this number (packs when sold in packs). Empty means 0."),
    "target_stock": (14, QUANTITY, "Optional, stock items. How many you like to have after restocking. Not lower than minimum_stock. Empty means 0."),
    "opening_quantity": (16, QUANTITY, "Optional: stock you have now, e.g. 10 cables. When sold in packs, count the PIECES (500 sheets = "
                                       "250 packs of 2). It is added as stock you own, not a purchase. Must be empty for SERVICE rows."),
    "opening_unit_cost": (16, MONEY, "Buying price for ONE item (for ONE piece when sold in packs). Opening stock value = opening_quantity "
                                     "× opening_unit_cost. If you do not know it, leave it empty: the stock is marked cost pending and "
                                     "you add the price later."),
    "preferred_shelf": (18, TEXT, "Optional: the existing compartment code or label where this item is kept, from the list. Empty keeps "
                                  "the stock as not yet placed. Must be empty for SERVICE rows. Shelves are never created by an import."),
}
assert list(FIELDS) == COLUMNS
INK, SOFT, REQUIRED_FILL, OPTIONAL_FILL = "1F2937", "F3F4F6", "FDE68A", "DBEAFE"
# One value per COLUMNS heading, in order.
SAMPLE_ROWS = [
    ["Oraimo Type-C cable", "Mobile Accessories > Cables", "Oraimo", "Braided fast-charging cable", "STOCK_ITEM", "1m", "", "",
     "ORAIMO-TCC-1M", "6974415130015", 300, 250, "", 3, 15, 10, 180, "AL1C1S1"],
    ["Oraimo Type-C cable", "Mobile Accessories > Cables", "Oraimo", "Braided fast-charging cable", "STOCK_ITEM", "2m", "", "",
     "", "", 400, 350, "", 2, 10, 6, 240, "AL1C1S1"],
    ["Amaya fast wall charger", "Mobile Accessories > Chargers > Wall Chargers", "Amaya", "", "", "", "", "", "AMAYA-FWC", "", 350, 300,
     "", 2, 6, 4, 160, "AL1C2S1"],
    ["Cable clips", "Mobile Accessories > Cable Organisers", "", "Adhesive clips", "STOCK_ITEM", "", 10, "clip", "", "", 20, "",
     "", 5, 20, 200, 1.5, "AL1C1S1"],
    ["Photocopy A4 black & white", "Services > Printing", "", "Per page", "SERVICE", "", "", "", "SRV-COPY-A4", "", 10, 10, 3,
     "", "", "", "", ""],
]
assert all(len(row) == len(COLUMNS) for row in SAMPLE_ROWS)


def _style_header(cell, required):
    cell.font = Font(bold=True, color=INK)
    cell.fill = PatternFill("solid", fgColor=REQUIRED_FILL if required else OPTIONAL_FILL)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = Border(bottom=Side(style="medium", color=INK))


def business_lists():
    """This business's category paths, brands and exact compartments (code, description) for the dropdowns."""
    from apps.inventory.models import Shelf
    from .models import Brand, Category
    categories = {c.pk: c for c in Category.objects.filter(active=True)}

    def path(c):
        parts = []
        while c:
            parts.append(c.name)
            c = categories.get(c.parent_id)
        return " > ".join(reversed(parts))
    from .models import shelves_in_use
    shelves = Shelf.objects.filter(active=True, is_virtual=False, zone__active=True, level__active=True, level__stack__active=True)
    return {
        "uses_shelves": shelves_in_use(),
        "categories": sorted((path(c) for c in categories.values()), key=str.casefold),
        "brands": sorted(Brand.objects.filter(active=True).values_list("name", flat=True), key=str.casefold),
        "shelves": [(s.physical_label or s.code, f"{s.level.stack.display_name}, level {s.level.level_number}")
                    for s in shelves.select_related("level__stack").order_by("level__stack__display_name", "level__level_number", "position_in_level", "code")],
    }


def build_template(lists=None, sample=False):
    """The .xlsx template. lists defaults to this business's; sample=True also fills the Products sheet with SAMPLE_ROWS."""
    lists = lists if lists is not None else business_lists()
    headings = [h for h in COLUMNS if h != "preferred_shelf" or lists.get("uses_shelves", True)]
    samples = [[row[COLUMNS.index(h)] for h in headings] for row in SAMPLE_ROWS]
    index = {heading: i for i, heading in enumerate(headings, start=1)}
    last = get_column_letter(len(headings))
    book = Workbook()
    sheet = book.active
    sheet.title = "Products"
    for col, heading in enumerate(headings, start=1):
        width, fmt, help_text = FIELDS[heading]
        cell = sheet.cell(row=1, column=col, value=heading)
        _style_header(cell, heading in REQUIRED)
        cell.comment = Comment(help_text, "LinTech", width=320, height=160)
        sheet.column_dimensions[get_column_letter(col)].width = width
        for row in range(2, ROWS + 2):
            sheet.cell(row=row, column=col).number_format = fmt
    if sample:
        for r, values in enumerate(samples, start=2):
            for col, value in enumerate(values, start=1):
                sheet.cell(row=r, column=col, value=value if value != "" else None)
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "B2"
    # An Excel table: banded rows, a filter on every heading, and new rows typed under it take its formats along.
    table = Table(displayName="Products", ref=f"A1:{last}{ROWS + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    sheet.add_table(table)

    # Hidden lists behind the dropdowns.
    hidden = book.create_sheet("Lists")
    columns = {"Types": ["STOCK_ITEM", "SERVICE"], "Categories": lists["categories"], "Brands": lists["brands"],
               "Shelves": [code for code, _ in lists["shelves"]] if "preferred_shelf" in index else []}
    ranges = {}
    for col, (title, values) in enumerate(columns.items(), start=1):
        hidden.cell(row=1, column=col, value=title).font = Font(bold=True)
        for row, value in enumerate(values, start=2):
            hidden.cell(row=row, column=col, value=value)
        letter = get_column_letter(col)
        ranges[title] = f"Lists!${letter}$2:${letter}${max(len(values), 1) + 1}"
    hidden.sheet_state = "hidden"

    def rule(heading, validation):
        letter = get_column_letter(index[heading])
        validation.add(f"{letter}2:{letter}{ROWS + 1}")
        sheet.add_data_validation(validation)

    def dropdown(heading, source, strict, message):
        rule(heading, DataValidation(type="list", formula1=source, allow_blank=True, showErrorMessage=True,
                                     errorStyle="stop" if strict else "information", error=message, errorTitle="Check this value"))

    dropdown("product_type", ranges["Types"], True, "Choose STOCK_ITEM or SERVICE, or leave it empty.")
    if lists["categories"]:
        dropdown("category", ranges["Categories"], False, "This category is new: it will be created when you import, "
                                                          "every level of the path exactly where it is written.")
    if lists["brands"]:
        dropdown("brand", ranges["Brands"], False, "This brand is new: it will be created when you import.")
    if lists["shelves"] and "preferred_shelf" in index:
        dropdown("preferred_shelf", ranges["Shelves"], True, "Choose an existing compartment from the list, or leave it empty.")
    rule("pack_size", DataValidation(type="whole", operator="between", formula1="1", formula2="100000", allow_blank=True,
                                     showErrorMessage=True, errorTitle="Whole number",
                                     error="Pieces in one pack: a whole number from 2 to 100,000. Empty or 1 means sold one at a time."))
    for heading in ("selling_price", "minimum_price", "service_cost", "minimum_stock", "target_stock", "opening_quantity", "opening_unit_cost"):
        rule(heading, DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True,
                                     error="Enter a number of 0 or more (no text, no negative numbers).", errorTitle="Numbers only"))

    # Instructions.
    guide = book.create_sheet("How to fill")
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 12
    guide.column_dimensions["C"].width = 100
    guide["A1"] = "How to fill the Products sheet"
    guide["A1"].font = Font(bold=True, size=14, color=INK)
    steps = [
        "1. Fill one row per sellable product or variant on the 'Products' sheet, just below the headings. Do not rename, move or "
        "add headings in row 1. Keep this file: add new rows at the bottom whenever you have more products and upload it again.",
        "2. Yellow columns are required. Blue columns are optional. Hover over a heading to see its note.",
        "3. Each row is saved exactly as the Add Product form would save it, with the same defaults and checks.",
        "4. Save the file as .xlsx and upload it in the app: Products > Import spreadsheet.",
        "5. The app checks every row first and shows every mistake with its row and column. Nothing is saved until you confirm, "
        "and then all rows are saved together or none.",
        "6. Categories: write the full path from the top, each level separated by > : Parent > Child > Grandchild > ... "
        "(as many levels as you need).",
        "   e.g. Mobile Accessories > Chargers > Car Chargers. The first name is always a top-level category. Every level not yet in "
        "the system is created, parent first, exactly where it is written; existing levels are reused (case and spaces ignored).",
        "7. Packs: put the pieces in one pack in pack_size (e.g. 10) and what a piece is called in piece_name (e.g. clip). The price "
        "is for one pack; opening stock is counted in pieces with the buying price of one piece, as in the form.",
        "8. Opening stock is stock you already own: it is not recorded as a purchase or money owed. Without a buying price it is "
        "marked cost pending until you add the price.",
        "9. Leave sku empty to use the code the form would suggest. Imported products are for stock management only: they are not "
        "visible or orderable online, so no photos are needed.",
        "10. Uploading the same file again is safe: rows already imported are skipped, never doubled. To change an existing "
        "product, edit it on its product page.",
    ]
    for row, text in enumerate(steps, start=3):
        guide.cell(row=row, column=1, value=text)
    start = len(steps) + 5
    for col, title in enumerate(["Column", "Required?", "What to write"], start=1):
        _style_header(guide.cell(row=start, column=col, value=title), False)
    for offset, heading in enumerate(headings, start=1):
        guide.cell(row=start + offset, column=1, value=heading).font = Font(bold=True)
        guide.cell(row=start + offset, column=2, value="Yes" if heading in REQUIRED else "No")
        note = guide.cell(row=start + offset, column=3, value=FIELDS[heading][2])
        note.alignment = Alignment(wrap_text=True, vertical="top")
        if offset % 2:
            for col in (1, 2, 3):
                guide.cell(row=start + offset, column=col).fill = PatternFill("solid", fgColor=SOFT)
    if lists["shelves"] and "preferred_shelf" in index:
        row = start + len(headings) + 2
        guide.cell(row=row, column=1, value="Your shelf compartments").font = Font(bold=True, size=12)
        for offset, (code, name) in enumerate(lists["shelves"], start=1):
            guide.cell(row=row + offset, column=1, value=code)
            guide.cell(row=row + offset, column=3, value=name)

    # A worked example, on its own sheet so it is never imported by mistake.
    example = book.create_sheet("Example")
    for col, heading in enumerate(headings, start=1):
        _style_header(example.cell(row=1, column=col, value=heading), heading in REQUIRED)
        example.column_dimensions[get_column_letter(col)].width = FIELDS[heading][0]
        for r, values in enumerate(samples, start=2):
            cell = example.cell(row=r, column=col, value=values[col - 1] if values[col - 1] != "" else None)
            cell.number_format = FIELDS[heading][1]
    example.cell(row=len(samples) + 3, column=1,
                 value="Example only: this sheet is not imported. Fill the 'Products' sheet.").font = Font(italic=True, color="B91C1C")
    example.freeze_panes = "B2"

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
