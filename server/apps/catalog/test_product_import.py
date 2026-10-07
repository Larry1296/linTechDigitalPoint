import io
import json
from decimal import Decimal

import pytest
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.accounting.models import JournalEntry
from apps.catalog.models import Brand, Category, Product, ProductImport, ProductVariant
from apps.catalog.product_import import COLUMNS
from apps.core.models import AuditLog
from apps.inventory.models import Movement, Shelf, ShelfLevel, ShelfStack, ShelfStackCategory, StockBalance, StockLot, VariantPreferredLocation, Zone

URL = "/api/v1/catalog/products/import/"


def workbook(*rows, headings=COLUMNS):
    """An .xlsx with the template headings; each row is a dict of column → value."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Products"
    sheet.append(list(headings))
    for row in rows:
        sheet.append([row.get(h) for h in headings])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def post(client, content, commit=False, racks=None, name="products.xlsx"):
    data = {"file": SimpleUploadedFile(name, content), "commit": "true" if commit else "false"}
    if racks is not None:
        data["rack_choices"] = json.dumps(racks)
    return client.post(URL, data, format="multipart")


def errors(body):
    return {row["row"]: [(e["column"], e["message"]) for e in row["errors"]] for row in body["rows"] if row["errors"]}


def nothing_written():
    return not (Product.objects.exists() or ProductVariant.objects.exists() or Brand.objects.exists() or StockLot.objects.exists()
                or JournalEntry.objects.exists() or ProductImport.objects.exists() or Category.objects.exclude(name="Mobile Accessories").exists())


@pytest.fixture(autouse=True)
def branch(db):
    from apps.core.models import Store
    return Store.objects.create(name="Main shop")


@pytest.fixture
def client():
    owner = User.objects.create_superuser("import-owner", "import@example.test", "Strong-pass-1296")
    api = APIClient()
    api.force_authenticate(owner)
    return api


def compartment(zone, stack, level, label):
    return Shelf.objects.create(zone=zone, level=level, code=f"{zone.code}-{label}", physical_label=label, display_name=label, x=0, y=0, width=1, height=1)


@pytest.fixture
def shop(branch):
    """A Mobile Accessories rack with two exact compartments, as in the real shop."""
    zone = Zone.objects.create(store=branch, code="A", name="Zone A")
    rack = ShelfStack.objects.create(zone=zone, code="A-R1", display_name="Mobile Accessories Rack", x=0, y=0, width=1, height=1, depth=1, number_of_levels=1)
    level = ShelfLevel.objects.create(stack=rack, level_number=1, y_position=0, height=1)
    root = Category.objects.create(name="Mobile Accessories", slug="mobile-accessories")
    ShelfStackCategory.objects.create(stack=rack, category=root, is_primary=True)
    return {"zone": zone, "rack": rack, "level": level, "root": root,
            "shelf": compartment(zone, rack, level, "AL1C1S1"), "other_shelf": compartment(zone, rack, level, "AL1C2S1")}


CABLE = {"name": "Oraimo Type-C cable", "category": "Mobile Accessories > Cables", "brand": "Oraimo", "variant_name": "1m", "sku": "ORAIMO-TCC-1M",
         "barcode": "6974415130015", "selling_price": 300, "minimum_price": 250, "minimum_stock": 3, "target_stock": 15, "opening_quantity": 10,
         "opening_unit_cost": 180, "preferred_shelf": "AL1C1S1"}
SERVICE = {"name": "Photocopy A4", "category": "Mobile Accessories > Services", "product_type": "SERVICE", "sku": "SRV-COPY", "selling_price": 10}


# Template

def test_template_has_exactly_the_documented_headings_and_this_shops_lists(client, shop):
    response = client.get("/api/v1/catalog/products/import-template/")
    assert response.status_code == 200 and response["Content-Disposition"].endswith('products-template.xlsx"')
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames[0] == "Products"
    assert [c.value for c in book["Products"][1]] == COLUMNS == [
        "name", "category", "brand", "description", "product_type", "variant_name", "pack_size", "piece_name", "sku", "barcode",
        "selling_price", "minimum_price", "service_cost", "minimum_stock", "target_stock", "opening_quantity", "opening_unit_cost",
        "preferred_shelf"]
    assert "Mobile Accessories" in [c.value for c in book["Lists"]["B"]] and "AL1C1S1" in [c.value for c in book["Lists"]["D"]]
    assert book["Products"].max_row >= 1 and book["Products"]["A2"].value is None
    # An Excel table, so rows added under the last one keep the dropdowns and formats.
    assert list(book["Products"].tables) == ["Products"]


def test_a_shop_without_shelves_gets_no_shelf_column(client, shop):
    from apps.tenancy.context import organization_id
    from apps.tenancy.models import Organization
    organization = Organization.objects.get(pk=organization_id())
    organization.enabled_modules = [m for m in organization.enabled_modules if m != "locations"]
    organization.save()
    book = load_workbook(io.BytesIO(client.get("/api/v1/catalog/products/import-template/").content))
    assert [c.value for c in book["Products"][1]] == [c for c in COLUMNS if c != "preferred_shelf"]


def test_filled_template_imports(client, shop):
    book = load_workbook(io.BytesIO(client.get("/api/v1/catalog/products/import-template/").content))
    sheet = book["Products"]
    for col, heading in enumerate(COLUMNS, start=1):
        sheet.cell(row=2, column=col, value=CABLE.get(heading))
    out = io.BytesIO()
    book.save(out)
    body = post(client, out.getvalue(), commit=True).json()
    assert body["committed"] is True, errors(body)
    assert ProductVariant.objects.get(sku="ORAIMO-TCC-1M").barcode == "6974415130015"


def test_online_columns_are_not_part_of_the_template(client, shop):
    response = post(client, workbook(CABLE, headings=[*COLUMNS, "ecommerce_visible", "online_orderable"]))
    assert response.status_code == 400
    assert "ecommerce_visible, online_orderable" in response.json()["file"] and "offline" in response.json()["file"]


def test_wrong_headings_are_rejected_with_the_expected_list(client):
    response = post(client, workbook({"name": "Pen"}, headings=["name", "category", "price", "sku"]))
    assert response.status_code == 400
    message = response.json()["file"]
    assert "unknown heading: price" in message and "missing: selling_price" in message and "opening_unit_cost" in message


# Offline, opening stock accounting

def test_imported_products_are_offline_and_opening_stock_goes_through_the_inventory_service(client, shop):
    preview = post(client, workbook(CABLE)).json()
    assert preview["committed"] is False and not errors(preview)
    assert preview["summary"]["opening_value"] == "1800.00" and preview["summary"]["total_quantity"] == "10"
    assert preview["summary"]["new_categories"] == ["Mobile Accessories > Cables"] and preview["summary"]["new_brands"] == ["Oraimo"]
    assert preview["rows"][0]["opening_value"] == "1800.00" and preview["rows"][0]["shelf"].startswith("AL1C1S1")
    assert nothing_written()

    response = post(client, workbook(CABLE), commit=True)
    assert response.status_code == 201, response.json()
    reference = response.json()["reference"]
    variant = ProductVariant.objects.get(sku="ORAIMO-TCC-1M")
    product = variant.product
    assert (product.ecommerce_visible, product.online_orderable) == (False, False)
    assert not product.images.exists()  # no photo needed
    assert (variant.name, variant.selling_price, variant.minimum_price, variant.minimum_stock, variant.target_stock) == ("1m", 300, 250, 3, 15)
    assert product.category.name == "Cables" and product.category.parent == shop["root"]
    # Quantity and cost are per sellable unit: 10 units at 180 each.
    lot = StockLot.objects.get(variant=variant)
    assert (lot.received_quantity, lot.remaining_quantity, lot.unit_cost, lot.cost_status, lot.reference) == (10, 10, 180, StockLot.COSTED, reference)
    assert StockBalance.objects.get(lot=lot).shelf == shop["shelf"] and StockBalance.objects.get(lot=lot).quantity == 10
    assert Movement.objects.get(lot=lot).movement_type == "OPENING_STOCK"
    assert VariantPreferredLocation.objects.get(variant=variant).shelf == shop["shelf"]
    # Owned stock is opening equity, not money owed to a supplier.
    entry = JournalEntry.objects.get(source_key=f"opening-stock:{lot.pk}")
    assert {(line.account.code, line.debit, line.credit) for line in entry.lines.all()} == {("1200", Decimal("1800.00"), 0), ("3000", 0, Decimal("1800.00"))}
    assert AuditLog.objects.filter(action="PRODUCT_CREATED", object_id=str(product.pk)).exists()
    assert AuditLog.objects.get(action="PRODUCTS_IMPORTED").object_id == reference
    assert ProductImport.objects.get().rows.get().variant == variant


def test_normal_product_entry_still_defaults_to_online(client, shop):
    response = client.post("/api/v1/catalog/products/create-with-stock/", {"name": "Earbuds", "category_id": shop["root"].pk, "sku": "EAR-1",
                                                                            "selling_price": "900", "ecommerce_visible": "true", "online_orderable": "true"})
    assert response.status_code == 201
    assert Product.objects.get(name="Earbuds").ecommerce_visible is True


def test_opening_stock_without_a_price_or_shelf_is_kept_as_the_form_keeps_it(client, shop):
    """As in the Add Product form: no buying price marks the stock cost pending (never valued at 0); no shelf keeps it unplaced."""
    rows = [
        {**CABLE, "sku": "A", "barcode": None, "opening_unit_cost": None},
        {**CABLE, "sku": "B", "barcode": None, "opening_unit_cost": 0},
        {**CABLE, "sku": "C", "barcode": None, "preferred_shelf": None},
    ]
    preview = post(client, workbook(*rows)).json()
    assert not errors(preview)
    assert any("cost pending" in note for note in preview["rows"][0]["notes"])
    assert any("not yet placed" in note for note in preview["rows"][2]["notes"])
    assert post(client, workbook(*rows), commit=True).json()["committed"] is True
    assert StockLot.objects.get(variant__sku="A").cost_status == StockLot.UNCOSTED
    assert StockLot.objects.get(variant__sku="B").cost_status == StockLot.UNCOSTED
    assert StockLot.objects.get(variant__sku="C").unit_cost == 180
    assert StockBalance.objects.get(lot__variant__sku="C").shelf.is_virtual


def test_opening_stock_and_prices_reject_negative_and_wrong_values(client, shop):
    rows = [
        {**CABLE, "sku": "D", "barcode": None, "opening_quantity": -2},
        {**CABLE, "sku": "E", "barcode": None, "opening_unit_cost": -5},
        {**CABLE, "sku": "F", "barcode": None, "selling_price": None, "minimum_stock": 5, "target_stock": 2},
        {**CABLE, "sku": "G", "barcode": None, "opening_quantity": "ten", "selling_price": 12.345},
        {**CABLE, "sku": "H", "barcode": None, "minimum_price": 301},
        {**CABLE, "sku": "I", "barcode": None, "minimum_price": -1},
    ]
    body = post(client, workbook(*rows), commit=True).json()
    assert body["committed"] is False
    found = errors(body)
    assert ("opening_quantity", "cannot be negative.") in found[2]
    assert ("opening_unit_cost", "cannot be negative.") in found[3]
    assert ("selling_price", "selling_price is required.") in found[4]
    assert ("target_stock", "(2) cannot be lower than minimum_stock (5).") in found[4]
    assert ("opening_quantity", "'ten' is not a number.") in found[5] and ("selling_price", "use at most 2 decimal places.") in found[5]
    assert ("minimum_price", "(301) cannot be higher than selling_price (300).") in found[6]
    assert ("minimum_price", "cannot be negative.") in found[7]
    assert nothing_written()


def test_rows_without_opening_stock_use_defaults(client, shop):
    row = {"name": "Screen protector", "category": "Mobile Accessories", "sku": "SP-1", "selling_price": 200}
    assert post(client, workbook(row), commit=True).status_code == 201
    variant = ProductVariant.objects.get(sku="SP-1")
    assert variant.minimum_price == 0  # no bargaining without a manager
    assert (variant.product.product_type, variant.name, variant.minimum_stock, variant.target_stock, variant.barcode) == ("STOCK_ITEM", "Standard", 0, 0, None)
    assert variant.product.brand is None and not StockLot.objects.filter(variant=variant).exists()
    assert not VariantPreferredLocation.objects.filter(variant=variant).exists()


def test_service_rows_create_no_stock_and_need_no_shelf(client, shop):
    assert post(client, workbook({**SERVICE, "category": "Mobile Accessories"}), commit=True).status_code == 201
    service = ProductVariant.objects.get(sku="SRV-COPY")
    assert service.product.product_type == "SERVICE" and not StockLot.objects.filter(variant=service).exists()
    bad = post(client, workbook({**SERVICE, "sku": "SRV-2", "category": "Mobile Accessories", "opening_quantity": 3,
                                 "opening_unit_cost": 5, "preferred_shelf": "AL1C1S1"})).json()
    assert {column for column, _ in errors(bad)[2]} == {"opening_quantity", "opening_unit_cost", "preferred_shelf"}


# Shelves

def test_preferred_shelf_must_be_one_existing_exact_compartment(client, shop, branch):
    zone_b = Zone.objects.create(store=branch, code="B", name="Zone B")
    rack_b = ShelfStack.objects.create(zone=zone_b, code="B-R1", display_name="Back Rack", x=0, y=0, width=1, height=1, depth=1, number_of_levels=1)
    ShelfStackCategory.objects.create(stack=rack_b, category=shop["root"], is_primary=True)
    compartment(zone_b, rack_b, ShelfLevel.objects.create(stack=rack_b, level_number=1, y_position=0, height=1), "AL1C1S1")  # same label as Zone A
    Shelf.objects.create(zone=shop["zone"], code="FLOOR-1", display_name="Floor pile", x=0, y=0, width=1, height=1)  # no rack level
    rows = [{**CABLE, "sku": s, "barcode": None, "preferred_shelf": shelf} for s, shelf in
            (("A", "NOPE-9"), ("B", "al1c1s1"), ("C", "FLOOR-1"), ("D", "A-AL1C2S1"), ("E", " al1c2s1 "))]
    shelves_before = Shelf.objects.count()
    body = post(client, workbook(*rows)).json()
    found = errors(body)
    assert "is not an existing shelf compartment" in found[2][0][1] and "never created" in found[2][0][1]
    assert found[3][0][0] == "preferred_shelf" and "ambiguous" in found[3][0][1] and "Back Rack" in found[3][0][1]
    assert "not an exact compartment" in found[4][0][1]
    assert 5 not in found and 6 not in found  # by code, and by label ignoring case and spaces
    assert Shelf.objects.count() == shelves_before


def test_compartment_on_a_rack_for_another_category_is_rejected(client, shop, branch):
    other = ShelfStack.objects.create(zone=shop["zone"], code="A-R2", display_name="Food Rack", x=0, y=0, width=1, height=1, depth=1, number_of_levels=1)
    ShelfStackCategory.objects.create(stack=other, category=Category.objects.create(name="Food", slug="food"), is_primary=True)
    compartment(shop["zone"], other, ShelfLevel.objects.create(stack=other, level_number=1, y_position=0, height=1), "FL1C1")
    body = post(client, workbook({**CABLE, "preferred_shelf": "FL1C1"})).json()
    (column, message), = errors(body)[2]
    assert column == "preferred_shelf" and "Food Rack" in message and "Mobile Accessories > Cables" in message


# Categories

def test_category_paths_create_parents_in_order_and_reuse_existing(client, shop):
    rows = [
        {**CABLE, "category": "  mobile ACCESSORIES >  Chargers ", "sku": "CHG-1", "barcode": None},   # existing root, new child
        {**CABLE, "category": "Mobile Accessories > chargers", "sku": "CHG-2", "barcode": None},        # same new child, other case
        {**CABLE, "category": "Mobile Accessories > Audio > Earbuds", "sku": "EAR-1", "barcode": None},  # new parent and child
        {**CABLE, "category": "Mobile Accessories > Audio", "sku": "AUD-1", "barcode": None},
    ]
    preview = post(client, workbook(*rows)).json()
    assert preview["summary"]["new_categories"] == ["Mobile Accessories > Chargers", "Mobile Accessories > Audio", "Mobile Accessories > Audio > Earbuds"]
    assert Category.objects.count() == 1
    assert post(client, workbook(*rows), commit=True).status_code == 201
    chargers = Category.objects.get(name="Chargers")
    assert chargers.parent == shop["root"] and Category.objects.filter(name__iexact="chargers").count() == 1
    audio = Category.objects.get(name="Audio")
    assert Category.objects.get(name="Earbuds").parent == audio and audio.parent == shop["root"]
    assert Category.objects.filter(name="Mobile Accessories").count() == 1
    assert {v.product.category for v in ProductVariant.objects.filter(sku__in=["CHG-1", "CHG-2"])} == {chargers}
    # Uploading more rows later reuses the categories the first import created.
    assert post(client, workbook({**CABLE, "category": "MOBILE ACCESSORIES > CHARGERS", "sku": "CHG-3", "barcode": None}), commit=True).status_code == 201
    assert Category.objects.count() == 4


def tree_path(category):
    parts = []
    while category:
        parts.append(category.name)
        category = category.parent
    return " > ".join(reversed(parts))


def test_deep_paths_create_every_missing_level_in_order(client, shop):
    audio = Category.objects.create(name="Audio", slug="audio", parent=shop["root"])
    paths = {
        "USB-C": "Mobile Accessories > Chargers > Car Chargers > USB-C",       # existing root, three new levels
        "WALL": " mobile accessories>chargers >  WALL chargers",               # reuses new levels, adds one
        "CHG": "Mobile Accessories > Chargers",                                # a new level used directly
        "WIFI": "Gadgets > Smart Home > Lighting > Bulbs > Wi-Fi",             # five levels, all new
        "EAR": "Mobile Accessories > Audio > Earbuds > Wireless",              # existing middle level, two new below
    }
    rows = [{**CABLE, "sku": sku, "barcode": None, "category": path} for sku, path in paths.items()]
    racks = {"Gadgets": shop["rack"].pk}
    preview = post(client, workbook(*rows), racks=racks).json()
    assert not errors(preview)
    # Parents always come before their children.
    assert preview["summary"]["new_categories"] == [
        "Mobile Accessories > Chargers", "Mobile Accessories > Chargers > Car Chargers", "Mobile Accessories > Chargers > Car Chargers > USB-C",
        "Mobile Accessories > Chargers > WALL chargers",
        "Gadgets", "Gadgets > Smart Home", "Gadgets > Smart Home > Lighting", "Gadgets > Smart Home > Lighting > Bulbs",
        "Gadgets > Smart Home > Lighting > Bulbs > Wi-Fi",
        "Mobile Accessories > Audio > Earbuds", "Mobile Accessories > Audio > Earbuds > Wireless"]
    assert Category.objects.count() == 2

    assert post(client, workbook(*rows), commit=True, racks=racks).status_code == 201
    assert Category.objects.count() == 2 + 11
    for sku, path in paths.items():
        category = ProductVariant.objects.get(sku=sku).product.category
        assert tree_path(category).casefold() == " > ".join(" ".join(p.split()) for p in path.split(">")).casefold()
    assert Category.objects.get(name="Earbuds").parent == audio
    assert Category.objects.filter(name="Chargers").count() == 1 and Category.objects.filter(name="Gadgets", parent=None).count() == 1
    assert ShelfStackCategory.objects.filter(category__name="Gadgets", stack=shop["rack"]).exists()

    # A later file reuses every existing level and adds only what is missing.
    more = {**CABLE, "sku": "ZIG", "barcode": None, "category": "GADGETS > smart home > Lighting > Bulbs > Zigbee"}
    assert post(client, workbook(more), commit=True).status_code == 201
    zigbee = Category.objects.get(name="Zigbee")
    assert tree_path(zigbee) == "Gadgets > Smart Home > Lighting > Bulbs > Zigbee" and Category.objects.count() == 14


def test_a_path_always_starts_at_the_top_level(client, shop):
    nested = Category.objects.create(name="Chargers", slug="chargers", parent=shop["root"])
    row = {**CABLE, "category": "Chargers > USB-C"}
    body = post(client, workbook(row), commit=True, racks={"Chargers": shop["rack"].pk}).json()
    assert body["committed"] is True, errors(body)
    assert "new top-level 'Chargers' will be created" in body["rows"][0]["notes"][0]
    assert "Mobile Accessories > Chargers" in body["rows"][0]["notes"][0]
    usb_c = ProductVariant.objects.get().product.category
    assert tree_path(usb_c) == "Chargers > USB-C" and usb_c.parent.parent is None
    assert not Category.objects.filter(parent=nested).exists()


def test_category_path_errors(client, shop):
    Category.objects.create(name="Hidden", slug="hidden", parent=shop["root"], active=False)
    rows = [{**CABLE, "sku": "X2", "barcode": None, "category": "Mobile Accessories > Hidden > Deeper"},
            {**CABLE, "sku": "X3", "barcode": None, "category": "Mobile Accessories >  > Cables"},
            {**CABLE, "sku": "X4", "barcode": None, "category": "Mobile Accessories > Cables >"}]
    found = errors(post(client, workbook(*rows)).json())
    assert "'Mobile Accessories > Hidden' is hidden" in found[2][0][1]
    assert "has an empty part" in found[3][0][1] and "has an empty part" in found[4][0][1]


def test_a_category_without_a_rack_is_reported_and_only_assigned_to_the_chosen_rack(client, shop):
    row = {**CABLE, "category": "Car Accessories > Chargers"}
    preview = post(client, workbook(row)).json()
    (column, message), = errors(preview)[2]
    assert column == "category" and "'Car Accessories > Chargers' is not on a rack yet" in message
    issue, = preview["summary"]["rack_issues"]
    # Asked once, for the top category: the sub-category and every future one under it use the same rack.
    assert issue == {"category": "Car Accessories", "new": True, "rows": [2], "suggested_racks": [shop["rack"].pk]}
    assert "Choose a rack for 'Car Accessories'" in message
    assert preview["summary"]["racks"] == [{"id": shop["rack"].pk, "name": "Mobile Accessories Rack", "area": "Zone A"}]
    # Never assigned silently: a plain commit fails and changes nothing.
    assert post(client, workbook(row), commit=True).json()["committed"] is False
    assert ShelfStackCategory.objects.count() == 1 and nothing_written()

    # The reviewer chooses a rack for the parent; the sub-category is stored on it.
    choice = {"car accessories": shop["rack"].pk}
    checked = post(client, workbook(row), racks=choice).json()
    assert not errors(checked) and checked["summary"]["rack_choices"][0]["category"] == "Car Accessories"
    assert post(client, workbook(row), commit=True, racks=choice).status_code == 201
    parent = Category.objects.get(name="Car Accessories")
    assert ShelfStackCategory.objects.filter(stack=shop["rack"], category=parent, is_primary=False).exists()
    assert AuditLog.objects.filter(action="RACK_CATEGORY_ASSIGNED").count() == 1


def test_rack_choices_need_shelf_permission_and_a_real_rack(client, shop):
    clerk = User.objects.create_user("clerk", password="Strong-pass-1296", is_staff=True)
    from django.contrib.auth.models import Permission
    clerk.user_permissions.add(*Permission.objects.filter(codename__in=["add_product", "view_product"], content_type__app_label="catalog"))
    api = APIClient()
    api.force_authenticate(User.objects.get(pk=clerk.pk))
    row = {**CABLE, "category": "Car Accessories"}
    denied = post(api, workbook(row), racks={"Car Accessories": shop["rack"].pk})
    assert denied.status_code == 400 and "Shelf management" in denied.json()["rack_choices"]
    body = post(client, workbook(row), racks={"Car Accessories": 999999, "Unknown": shop["rack"].pk}).json()
    assert body["summary"]["file_errors"] == ["Rack choice for 'Car Accessories': choose an active rack.",
                                              "Rack choice for 'Unknown': that category is not in this file."]


def test_without_the_shelves_tool_categories_need_no_rack(client, shop):
    from apps.tenancy.context import organization_id
    from apps.tenancy.models import Organization
    organization = Organization.objects.get(pk=organization_id())
    organization.enabled_modules = [m for m in organization.enabled_modules if m != "locations"]
    organization.save()
    assert post(client, workbook({**SERVICE, "category": "Services > Printing"}), commit=True).status_code == 201
    assert ShelfStackCategory.objects.count() == 1


# Duplicates and conflicts

def test_duplicates_in_the_file_are_errors_and_existing_codes_are_skipped(client, shop):
    existing = Product.objects.create(name="Old charger", slug="old", category=shop["root"])
    ProductVariant.objects.create(product=existing, sku="CHG-OLD", barcode="111", selling_price=500)
    rows = [
        {**CABLE, "sku": "DUP-1", "barcode": "222"},
        {**CABLE, "sku": "dup-1", "barcode": "333"},
        {**CABLE, "sku": "U-3", "barcode": "222"},
        {**CABLE, "sku": "chg-old", "barcode": None},
        {**CABLE, "sku": "U-5", "barcode": "111"},
        {**CABLE, "sku": None, "name": None, "category": None, "barcode": None},
    ]
    body = post(client, workbook(*rows), commit=True).json()
    assert body["committed"] is False
    found = errors(body)
    assert ("sku", "dup-1 is also used on row 3. Each sku must be unique in the file.") not in found[2]
    assert ("sku", "DUP-1 is also used on row 3. Each sku must be unique in the file.") in found[2]
    assert ("sku", "dup-1 is also used on row 2. Each sku must be unique in the file.") in found[3]
    assert ("barcode", "222 is also used on row 4. Each barcode must be unique in the file.") in found[2]
    assert 5 not in found and 6 not in found
    actions = {r["row"]: r for r in body["rows"]}
    assert actions[5]["action"] == "skip" and "SKU CHG-OLD already exists (Old charger, Standard)" in actions[5]["notes"][0]
    assert actions[6]["action"] == "skip" and "Barcode 111 already belongs to Old charger (SKU CHG-OLD)" in actions[6]["notes"][0]
    assert {c for c, _ in found[7]} >= {"name", "sku", "category"}
    assert body["summary"]["errors"] == 4 and body["summary"]["skip"] == 2
    assert Product.objects.count() == 1 and not StockLot.objects.exists() and not ProductImport.objects.exists()


# All or nothing

def test_a_failure_while_saving_rolls_back_every_row(client, shop, monkeypatch):
    from apps.inventory import services
    real, calls = services.create_product_with_opening_stock, []

    def failing(**kwargs):
        calls.append(kwargs)
        if len(calls) == 2:
            from rest_framework.exceptions import ValidationError
            raise ValidationError({"sku": "simulated failure"})
        return real(**kwargs)
    monkeypatch.setattr(services, "create_product_with_opening_stock", failing)
    rows = [{**CABLE, "category": "Mobile Accessories > New", "brand": "Brandnew"}, {**CABLE, "sku": "SECOND", "barcode": None}]
    response = post(client, workbook(*rows), commit=True)
    assert response.status_code == 400 and "Row 3 (SECOND) could not be saved" in response.json()["file"]
    assert "Nothing was imported" in response.json()["file"]
    assert nothing_written() and not Movement.objects.exists() and not AuditLog.objects.filter(action__in=["PRODUCTS_IMPORTED", "CATEGORY_CREATED"]).exists()


# Repeat uploads

def test_uploading_the_same_file_again_changes_nothing(client, shop):
    content = workbook(CABLE, {**SERVICE, "category": "Mobile Accessories"})
    first = post(client, content, commit=True).json()
    assert first["summary"]["create"] == 2
    again = post(client, content).json()
    assert again["summary"]["previous_import"]["reference"] == first["reference"]
    assert [r["action"] for r in again["rows"]] == ["skip", "skip"] and "Already imported in" in again["rows"][0]["notes"][0]
    committed = post(client, content, commit=True)
    assert committed.status_code == 200 and committed.json()["committed"] is True and committed.json()["reference"] == ""
    assert Product.objects.count() == 2 and StockLot.objects.count() == 1 and JournalEntry.objects.count() == 1
    assert ProductImport.objects.count() == 1


def test_a_grown_file_adds_only_new_rows_and_a_changed_row_is_skipped(client, shop):
    assert post(client, workbook(CABLE), commit=True).status_code == 201
    # The same row re-typed with different spacing, case and number formatting is still the same row.
    retyped = {**CABLE, "name": " oraimo  type-c CABLE", "selling_price": "300.00", "category": "mobile accessories>cables"}
    grown = workbook(retyped, {**CABLE, "sku": "ORAIMO-TCC-2M", "variant_name": "2m", "barcode": None})
    body = post(client, grown, commit=True).json()
    assert body["committed"] is True and body["summary"]["create"] == 1 and body["summary"]["skip"] == 1
    assert StockLot.objects.filter(variant__sku="ORAIMO-TCC-1M").count() == 1
    changed = post(client, workbook({**CABLE, "opening_quantity": 20}), commit=True).json()
    assert changed["committed"] is True and changed["summary"]["error_count"] == 0 and changed["summary"]["create"] == 0
    assert changed["rows"][0]["action"] == "skip" and "already exists" in changed["rows"][0]["notes"][0]
    assert StockLot.objects.filter(variant__sku="ORAIMO-TCC-1M").count() == 1


def test_products_already_in_the_system_are_skipped_and_the_rest_imported(client, shop):
    existing = Product.objects.create(name="Old charger", slug="old", category=shop["root"])
    ProductVariant.objects.create(product=existing, sku="CHG-OLD", barcode="111", selling_price=500)
    rows = [CABLE, {**CABLE, "sku": "chg-old", "barcode": None, "selling_price": 999},
            {**CABLE, "sku": "NEW-2", "barcode": "111"}, {**CABLE, "sku": "chg-old", "barcode": None}]
    body = post(client, workbook(*rows), commit=True).json()
    assert body["committed"] is True, errors(body)
    assert body["summary"]["create"] == 1 and body["summary"]["skip"] == 3
    assert ProductVariant.objects.get(sku="CHG-OLD").selling_price == 500
    assert not ProductVariant.objects.filter(sku="NEW-2").exists()


def test_import_needs_product_creation_permission(shop):
    clerk = User.objects.create_user("clerk", password="Strong-pass-1296", is_staff=True)
    api = APIClient()
    api.force_authenticate(clerk)
    assert post(api, workbook(CABLE)).status_code == 403


def test_csv_with_the_same_headings_still_works(client, shop):
    row = {"name": "Pen", "category": "Mobile Accessories", "sku": "PEN-1", "selling_price": "20"}
    text = ",".join(COLUMNS) + "\n" + ",".join(row.get(c, "") for c in COLUMNS) + "\n"
    response = client.post(URL, {"file": SimpleUploadedFile("p.csv", text.encode()), "commit": "true"}, format="multipart")
    assert response.status_code == 201, response.json()


def test_files_filled_from_the_earlier_template_still_import(client, shop):
    earlier = [c for c in COLUMNS if c not in {"pack_size", "piece_name", "service_cost"}]
    body = post(client, workbook(CABLE, headings=earlier), commit=True).json()
    assert body["committed"] is True, errors(body)
    assert post(client, workbook(CABLE, headings=earlier)).json()["rows"][0]["action"] == "skip"


# Saved exactly as the Add Product form saves it

def test_packs_are_priced_per_pack_and_opening_stock_is_counted_in_pieces(client, shop):
    clips = {"name": "Cable clips", "category": "Mobile Accessories > Cable Organisers", "pack_size": 10, "piece_name": "clips",
             "selling_price": 20, "minimum_stock": 5, "opening_quantity": 205, "opening_unit_cost": 1.5, "preferred_shelf": "AL1C1S1"}
    body = post(client, workbook(clips), commit=True).json()
    assert body["committed"] is True, errors(body)
    variant = ProductVariant.objects.get(product__name="Cable clips")
    # The form names a default variant after the pack and suggests the SKU from the name and variant.
    assert (variant.name, variant.sku, variant.pack_size, variant.piece_name) == ("Pack of 10 clips", "CC-PACK-OF-10", 10, "clips")
    lot = StockLot.objects.get(variant=variant)
    assert (lot.received_quantity, lot.unit_cost) == (Decimal("20.500"), Decimal("15.00"))
    assert Category.objects.get(name="Cable Organisers").parent == shop["root"]


def test_pack_and_service_columns_are_checked_like_the_form(client, shop):
    rows = [
        {**CABLE, "sku": "P1", "barcode": None, "pack_size": 2.5},
        {**CABLE, "sku": "P2", "barcode": None, "piece_name": "sheet"},
        {**SERVICE, "sku": "S1", "pack_size": 4},
        {**SERVICE, "sku": "S2", "barcode": "123"},
        {**CABLE, "sku": "P3", "barcode": None, "service_cost": 40},
    ]
    found = errors(post(client, workbook(*rows)).json())
    assert found[2][0][0] == "pack_size" and "whole number" in found[2][0][1]
    assert found[3] == [("piece_name", "is only used with a pack_size of 2 or more.")]
    assert found[4][0][0] == "pack_size"
    assert found[5][0][0] == "barcode"
    assert found[6] == [("service_cost", "is only for SERVICE rows. Leave it empty for stock items.")]


def test_service_cost_is_saved_on_a_service(client, shop, branch):
    ShelfStackCategory.objects.create(stack=shop["rack"], category=Category.objects.create(name="Services", slug="services"))
    service = {"name": "Windows installation", "category": "Services", "product_type": "SERVICE", "selling_price": 1500, "service_cost": 200}
    body = post(client, workbook(service), commit=True).json()
    assert body["committed"] is True, errors(body)
    variant = ProductVariant.objects.get(product__name="Windows installation")
    assert (variant.service_cost, variant.sku, variant.name) == (Decimal("200.00"), "WI-SERVICE", "Standard")


def test_an_empty_sku_is_the_one_the_form_suggests_and_a_reupload_skips_it(client, shop):
    from apps.catalog.product_import import suggested_sku
    assert suggested_sku("Oraimo Type-C cable", "Oraimo", "1m") == "ORAIMO-TCC-1M"
    assert suggested_sku("Original Normal Full Charger", "Oraimo") == "ORAIMO-ONFC"
    assert suggested_sku("Power Bank 20000", "") == "PB-20000"
    assert suggested_sku("C-C HEAD(45W)", "") == "CCH-45W"
    assert suggested_sku("EVERADY 2AA", "Everady", "EVERADY-2AA") == "EVERADY-2AA"
    assert suggested_sku("Wall charger", "", "Black") == "WC-BLACK"
    assert suggested_sku("Photocopy", "", "Standard", "SERVICE") == "PHOTOCOPY-SERVICE"
    # The same cases as the form's own tests (client/src/modules/products/sku.test.jsx), so both always agree.
    assert suggested_sku("Oraimo Original Normal Full Charger", "Oraimo") == "ORAIMO-ONFC"
    assert suggested_sku("Power Bank 20000", "Oraimo", "Black") == "ORAIMO-PB-20000-BLACK"
    assert suggested_sku("Smartphone Cover", "", "Samsung A15 / Black") == "SC-SAMSUNG-A15-BLACK"
    assert suggested_sku("  Café & cable! ") == "CC"
    assert suggested_sku("Router", "TP-Link") == "TP-LINK-ROUTER"
    assert suggested_sku("") == ""
    long = suggested_sku("Very long name " * 20, "A brand with a very long name indeed", "Blue with extra long variant description")
    assert len(long) <= 64 and long.startswith("A-BRAND-WITH")
    row = {**CABLE, "sku": None, "barcode": None}
    assert post(client, workbook(row), commit=True).json()["committed"] is True
    assert ProductVariant.objects.filter(sku="ORAIMO-TCC-1M").exists()
    assert post(client, workbook(row)).json()["rows"][0]["action"] == "skip"


def test_documented_sample_workbook_imports_cleanly(client, shop):
    from pathlib import Path
    from django.conf import settings
    content = (Path(settings.BASE_DIR).parent / "docs" / "products-import-sample.xlsx").read_bytes()
    preview = post(client, content).json()
    # Services > Printing has no rack yet: the review asks once for Services, and the reviewer chooses.
    assert [issue["category"] for issue in preview["summary"]["rack_issues"]] == ["Services"]
    body = post(client, content, commit=True, racks={"Services": shop["rack"].pk}).json()
    assert body["committed"] is True, errors(body)
    assert body["summary"]["create"] == 5 and body["summary"]["services"] == 1 and body["summary"]["opening_value"] == "4180.00"
    assert not Product.objects.filter(ecommerce_visible=True).exists()
